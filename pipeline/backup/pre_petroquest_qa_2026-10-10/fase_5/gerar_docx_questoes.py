"""Gera um .docx simples das questões, para revisão dos especialistas.

Só o essencial: enunciado, alternativas e a resposta correta. Sem nota, sem
dificuldade, sem similaridade — o objetivo é ler as questões, não auditar o
pipeline.

Uso, a partir da raiz do repositório:

    # o dataset filtrado inteiro (5.645 questões, documento longo)
    python pipeline/fase_5/gerar_docx_questoes.py

    # uma amostra de 5 questões por subtópico, bom para apresentação
    python pipeline/fase_5/gerar_docx_questoes.py --por-subtopico 5

    # sem a resposta junto da questão; gabarito numa tabela no fim
    python pipeline/fase_5/gerar_docx_questoes.py --por-subtopico 5 --gabarito-no-fim

Requer `python-docx` (pip install python-docx).
"""

from __future__ import annotations

import argparse
import json
import random
import re
from collections import OrderedDict, defaultdict
from datetime import date
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent

LETRAS = "ABCDEFGH"
PADRAO_ENTRADA = (AQUI / "saida_fase5" /
                  "questoes_filtradas_cota_cluster_piso0.45.jsonl")


# Caracteres de controle não são válidos em XML e derrubam o python-docx. No
# dataset atual aparece exatamente um (um \x00 no meio de "óleo-água", resíduo
# da geração); vira espaço e o espaço em excesso é colapsado.
RE_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def limpar(texto: str) -> str:
    return re.sub(r"\s+", " ", RE_CONTROLE.sub(" ", texto or "")).strip()


def marcar(paragrafo, nome: str, indice: int) -> None:
    """Põe um bookmark do Word no parágrafo, alvo dos links do sumário.

    O python-docx não expõe bookmark nem link interno, então os dois saem em
    XML na mão (`w:bookmarkStart`/`w:bookmarkEnd` e `w:hyperlink w:anchor`).
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    inicio = OxmlElement("w:bookmarkStart")
    inicio.set(qn("w:id"), str(indice))
    inicio.set(qn("w:name"), nome)
    fim = OxmlElement("w:bookmarkEnd")
    fim.set(qn("w:id"), str(indice))
    paragrafo._p.insert(0, inicio)
    paragrafo._p.append(fim)


def link_interno(paragrafo, texto: str, ancora: str, negrito: bool = False) -> None:
    """Texto clicável que salta para o bookmark `ancora`, no mesmo documento."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), ancora)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    cor = OxmlElement("w:color")
    cor.set(qn("w:val"), "1F4E79")
    sublinhado = OxmlElement("w:u")
    sublinhado.set(qn("w:val"), "single")
    props.append(cor)
    props.append(sublinhado)
    if negrito:
        props.append(OxmlElement("w:b"))
    run.append(props)
    elemento = OxmlElement("w:t")
    elemento.text = texto
    elemento.set(qn("xml:space"), "preserve")
    run.append(elemento)
    link.append(run)
    paragrafo._p.append(link)


def numero_de_pagina(paragrafo) -> None:
    """Campo PAGE do Word, para o rodapé."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            marca = OxmlElement("w:fldChar")
            marca.set(qn("w:fldCharType"), tipo)
            run._r.append(marca)
        else:
            instrucao = OxmlElement("w:instrText")
            instrucao.set(qn("xml:space"), "preserve")
            instrucao.text = texto
            run._r.append(instrucao)


def carregar(caminho: Path) -> list[dict]:
    return [json.loads(l) for l in caminho.read_text(encoding="utf-8").splitlines()
            if l.strip()]


def _cotas_proporcionais(tamanhos: dict[str, int], total: int) -> dict[str, int]:
    """Distribui `total` entre os subtópicos proporcionalmente ao que cada um tem.

    Maiores restos para desempatar, mínimo de 1 por subtópico com questão, e
    nunca mais do que o subtópico tem.
    """
    disponivel = {k: v for k, v in tamanhos.items() if v > 0}
    if not disponivel:
        return {}
    soma = sum(disponivel.values())
    if total >= soma:
        return dict(disponivel)

    exatas = {k: v * total / soma for k, v in disponivel.items()}
    cotas = {k: min(disponivel[k], max(1, int(v))) for k, v in exatas.items()}
    # ajusta a sobra (ou o excesso) pelos maiores restos
    resto = total - sum(cotas.values())
    ordem = sorted(disponivel, key=lambda k: exatas[k] - int(exatas[k]),
                   reverse=resto > 0)
    while resto != 0:
        mudou = False
        for k in ordem:
            if resto > 0 and cotas[k] < disponivel[k]:
                cotas[k] += 1; resto -= 1; mudou = True
            elif resto < 0 and cotas[k] > 1:
                cotas[k] -= 1; resto += 1; mudou = True
            if resto == 0:
                break
        if not mudou:
            break
    return cotas


def amostrar(questoes: list[dict], por_subtopico: int | None,
             por_topico: int | None, limite: int | None, seed: int) -> list[dict]:
    """Amostra por subtópico ou por tópico e/ou corta no total.

    `--por-topico N` reparte as N questões entre os subtópicos do tópico em
    proporção ao tamanho de cada um, com pelo menos uma de cada: o documento
    fica do tamanho pedido sem perder subtópico do sumário.
    """
    rng = random.Random(seed)
    por_chave: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for q in questoes:
        por_chave[(q["topico"], q["subtopico"])].append(q)

    if por_subtopico:
        escolhidas = []
        for grupo in por_chave.values():
            escolhidas.extend(rng.sample(grupo, min(por_subtopico, len(grupo))))
        questoes = escolhidas
    elif por_topico:
        escolhidas = []
        topicos: dict[str, list[str]] = defaultdict(list)
        for (topico, subtopico) in por_chave:
            topicos[topico].append(subtopico)
        for topico, subtopicos in topicos.items():
            tamanhos = {s: len(por_chave[(topico, s)]) for s in subtopicos}
            for subtopico, cota in _cotas_proporcionais(tamanhos, por_topico).items():
                grupo = por_chave[(topico, subtopico)]
                escolhidas.extend(rng.sample(grupo, min(cota, len(grupo))))
        questoes = escolhidas

    if limite:
        questoes = questoes[:limite]

    # devolve na ordem original do arquivo, para o agrupamento ficar estável
    escolhidos = {id(q) for q in questoes}
    return [q for chave in por_chave for q in por_chave[chave] if id(q) in escolhidos]


def agrupar(questoes: list[dict]) -> "OrderedDict[str, OrderedDict[str, list[dict]]]":
    """Tópico -> subtópico -> questões, preservando a ordem de primeira aparição."""
    arvore: OrderedDict[str, OrderedDict[str, list[dict]]] = OrderedDict()
    for q in questoes:
        arvore.setdefault(q["topico"], OrderedDict()).setdefault(
            q["subtopico"], []).append(q)
    return arvore


def montar(questoes: list[dict], destino: Path, titulo: str,
           gabarito_no_fim: bool) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    documento = Document()

    estilo = documento.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)

    # -------------------------------------------------------------- capa --
    capa = documento.add_heading(titulo, level=0)
    capa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = documento.add_paragraph(
        f"{len(questoes)} questões · gerado em {date.today():%d/%m/%Y}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    rodape = documento.sections[0].footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    numero_de_pagina(rodape)

    arvore = agrupar(questoes)
    gabarito: list[tuple[int, str]] = []
    numero = 0

    # ------------------------------------------------------------ sumário --
    # Ancoras: um bookmark por tópico e por subtópico; o sumário aponta para
    # eles, então clicar na linha salta para a seção (Ctrl+clique no Word).
    ancoras: dict[tuple[str, str | None], str] = {}
    for i, (topico, subtopicos) in enumerate(arvore.items()):
        ancoras[(topico, None)] = f"topico_{i}"
        for j, subtopico in enumerate(subtopicos):
            ancoras[(topico, subtopico)] = f"topico_{i}_sub_{j}"

    documento.add_heading("Sumário", level=1)
    intro = documento.add_paragraph(
        "Clique em um tópico ou subtópico para ir até ele "
        "(no Word, Ctrl+clique).")
    intro.runs[0].italic = True
    intro.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    for topico, subtopicos in arvore.items():
        n_topico = sum(len(g) for g in subtopicos.values())
        linha = documento.add_paragraph()
        linha.paragraph_format.space_after = Pt(2)
        link_interno(linha, topico, ancoras[(topico, None)], negrito=True)
        contagem = linha.add_run(f"  ·  {n_topico} questões")
        contagem.bold = True
        for subtopico, grupo in subtopicos.items():
            if subtopico == topico:
                continue
            sublinha = documento.add_paragraph()
            sublinha.paragraph_format.left_indent = Pt(18)
            sublinha.paragraph_format.space_after = Pt(0)
            link_interno(sublinha, subtopico, ancoras[(topico, subtopico)])
            sublinha.add_run(f"  ·  {len(grupo)}")

    total = documento.add_paragraph()
    total.paragraph_format.space_before = Pt(10)
    total.add_run(f"Total: {len(questoes)} questões").bold = True

    # -------------------------------------------------------------- corpo --
    indice_marcador = 1
    for topico, subtopicos in arvore.items():
        documento.add_page_break()
        cabecalho = documento.add_heading(topico, level=1)
        marcar(cabecalho, ancoras[(topico, None)], indice_marcador)
        indice_marcador += 1
        for subtopico, grupo in subtopicos.items():
            if subtopico != topico:
                sub_cabecalho = documento.add_heading(subtopico, level=2)
                marcar(sub_cabecalho, ancoras[(topico, subtopico)],
                       indice_marcador)
                indice_marcador += 1
            for questao in grupo:
                numero += 1
                correta = int(questao["correct_answer_index"])

                enunciado = documento.add_paragraph()
                enunciado.paragraph_format.space_before = Pt(10)
                enunciado.paragraph_format.space_after = Pt(4)
                rotulo = enunciado.add_run(f"{numero}. ")
                rotulo.bold = True
                enunciado.add_run(limpar(questao["stem"]))

                for i, alternativa in enumerate(questao["alternatives"]):
                    linha = documento.add_paragraph()
                    linha.paragraph_format.left_indent = Pt(18)
                    linha.paragraph_format.space_after = Pt(2)
                    marca = linha.add_run(f"{LETRAS[i]}) ")
                    texto = linha.add_run(limpar(alternativa))
                    if i == correta and not gabarito_no_fim:
                        marca.bold = True
                        texto.bold = True

                gabarito.append((numero, LETRAS[correta]))

    # ---------------------------------------------------------- gabarito --
    if gabarito_no_fim:
        documento.add_page_break()
        documento.add_heading("Gabarito", level=1)
        colunas = 6
        linhas = (len(gabarito) + colunas - 1) // colunas
        tabela = documento.add_table(rows=linhas, cols=colunas)
        tabela.style = "Table Grid"
        for posicao, (num, letra) in enumerate(gabarito):
            celula = tabela.cell(posicao % linhas, posicao // linhas)
            celula.text = f"{num}. {letra}"
    else:
        rodape = documento.add_paragraph()
        rodape.paragraph_format.space_before = Pt(14)
        nota = rodape.add_run("A alternativa em negrito é a resposta correta.")
        nota.italic = True
        nota.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.save(destino)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--questoes", type=Path, default=PADRAO_ENTRADA,
                        help="JSONL de entrada (padrão: o dataset filtrado da fase 5).")
    parser.add_argument("--saida", type=Path, default=None,
                        help="Caminho do .docx (padrão: entregaveis/Questoes_<n>.docx).")
    parser.add_argument("--titulo", default="Banco de Questões — FPSO",
                        help="Título da capa.")
    parser.add_argument("--por-subtopico", type=int, default=None,
                        help="Amostra N questões de cada subtópico.")
    parser.add_argument("--por-topico", type=int, default=None,
                        help="Amostra N questões de cada tópico, repartidas "
                             "entre os subtópicos em proporção ao tamanho.")
    parser.add_argument("--limite", type=int, default=None,
                        help="Teto de questões no documento.")
    parser.add_argument("--gabarito-no-fim", action="store_true",
                        help="Não marca a correta na questão; põe o gabarito no fim.")
    parser.add_argument("--seed", type=int, default=42,
                        help="Semente da amostragem (padrão 42).")
    args = parser.parse_args()

    questoes = carregar(args.questoes)
    print(f"{len(questoes)} questões em {args.questoes.name}")
    questoes = amostrar(questoes, args.por_subtopico, args.por_topico,
                        args.limite, args.seed)

    saida = args.saida or (AQUI / "entregaveis" /
                           f"Questoes_{len(questoes)}.docx")
    montar(questoes, saida, args.titulo, args.gabarito_no_fim)
    print(f"{len(questoes)} questões -> {saida}")


if __name__ == "__main__":
    main()
