"""Gera um caderno DOCX editável para a curadoria humana de questões da Fase 3."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
INPUT_PATH = ROOT / "pipeline" / "fase_3" / "saida_fase3" / "questoes_fase3.jsonl"
OUTPUT_DIR = ROOT / "pipeline" / "fase_4" / "entregaveis"
OUTPUT_PATH = OUTPUT_DIR / "Caderno_de_Curadoria_Questoes_Fase3.docx"

# Design preset: compact_reference_guide. Header pattern: memo_masthead.
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "5F6B7A"


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, color: str | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGINS_DXA.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_dxa / 1440)


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Larguras inválidas: {sum(widths_dxa)} DXA; esperado {CONTENT_WIDTH_DXA}.")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_dxa):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_border(cell, *, color: str = "D9DEE7", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, end])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("CADERNO DE CURADORIA | FASE 3")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("Página ")
    set_run_font(run, size=8.5, color=MUTED)
    add_page_field(footer)


def load_questions() -> list[dict]:
    questions = []
    with INPUT_PATH.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                questions.append(json.loads(line))
    return questions


def quality_key(question: dict) -> tuple:
    """Maior nota; em empates, melhor avaliação do juiz e menor soma de vícios."""
    judge = question.get("judge", {})
    judge_total = sum(float(judge.get(field, 0)) for field in (
        "correcao", "clareza", "alternativas", "relevancia",
    ))
    vicios_total = sum(float(value) for value in question.get("vicios", {}).values())
    return (-float(question.get("nota", 0)), -judge_total, vicios_total, question["id"])


def select_questions(questions: list[dict]) -> dict[str, list[dict]]:
    topics = sorted({question["topico"] for question in questions})
    selected = {}
    for topic in topics:
        ranked = sorted((q for q in questions if q["topico"] == topic), key=quality_key)
        if len(ranked) < 5:
            raise ValueError(f"O tópico {topic!r} possui apenas {len(ranked)} questões.")
        selected[topic] = ranked[:5]
    return selected


def add_title_page(doc: Document, selected: dict[str, list[dict]]) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    run = title.add_run("Caderno de Curadoria")
    set_run_font(run, size=24, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Questões selecionadas da Fase 3")
    set_run_font(run, size=14, color=MUTED)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    set_table_geometry(meta, [2700, 6660])
    for row, (label, value) in zip(meta.rows, (
        ("Seleção", f"5 questões de cada tópico, totalizando {sum(len(items) for items in selected.values())} questões."),
        ("Critério", "Maior nota final; em empate, melhor avaliação do juiz e menor soma de vícios."),
        ("Uso", "Preencher a decisão de curadoria e as observações para cada questão."),
    )):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_border(row.cells[0])
        set_cell_border(row.cells[1])
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(label), size=10.5, color=NAVY, bold=True)
        p = row.cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size=10.5)

    doc.add_paragraph()
    heading = doc.add_paragraph("Tópicos incluídos", style="Heading 2")
    heading.paragraph_format.space_before = Pt(14)
    topics = doc.add_table(rows=1, cols=2)
    topics.style = "Table Grid"
    set_table_geometry(topics, [7560, 1800])
    set_repeat_table_header(topics.rows[0])
    for cell, text in zip(topics.rows[0].cells, ("Tópico", "Questões")):
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=10, color=NAVY, bold=True)
    for topic, questions in selected.items():
        row = topics.add_row()
        for cell, text in zip(row.cells, (topic, str(len(questions)))):
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(text), size=9.5)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    note.paragraph_format.space_after = Pt(0)
    set_run_font(note.add_run("Orientação: "), size=10.5, color=NAVY, bold=True)
    set_run_font(note.add_run("marque uma decisão e registre observações objetivas sobre conteúdo, clareza, gabarito ou alternativas."), size=10.5, color=MUTED)


def add_metadata_table(doc: Document, question: dict) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [4680, 4680])
    set_repeat_table_header(table.rows[0])
    values = (
        ("Subtópico", question.get("subtopico", "-")),
        ("Faceta", question.get("faceta_titulo", "-")),
        ("ID", question["id"]),
        ("Nota de geração", f"{float(question.get('nota', 0)):.2f}"),
    )
    for cell, (label, value) in zip([c for row in table.rows for c in row.cells], values):
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(f"{label}: "), size=8.8, color=MUTED, bold=True)
        set_run_font(paragraph.add_run(str(value)), size=8.8, color="333333")


def add_question_form(doc: Document, question: dict, topic: str, index: int) -> None:
    doc.add_page_break()

    topic_line = doc.add_paragraph()
    topic_line.paragraph_format.space_after = Pt(3)
    set_run_font(topic_line.add_run(topic.upper()), size=9.5, color=DARK_BLUE, bold=True)

    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(7)
    set_run_font(heading.add_run(f"Questão {index} de 5"), size=13, color=BLUE, bold=True)
    add_metadata_table(doc, question)

    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(10)
    label.paragraph_format.space_after = Pt(3)
    set_run_font(label.add_run("ENUNCIADO"), size=9.5, color=MUTED, bold=True)
    stem = doc.add_paragraph()
    stem.paragraph_format.space_after = Pt(8)
    set_run_font(stem.add_run(question["stem"]), size=10.5)

    alternatives_label = doc.add_paragraph()
    alternatives_label.paragraph_format.space_after = Pt(3)
    set_run_font(alternatives_label.add_run("ALTERNATIVAS"), size=9.5, color=MUTED, bold=True)
    letters = "ABCDEF"
    for letter, alternative in zip(letters, question["alternatives"]):
        option = doc.add_paragraph()
        option.paragraph_format.left_indent = Inches(0.25)
        option.paragraph_format.first_line_indent = Inches(-0.25)
        option.paragraph_format.space_after = Pt(4)
        set_run_font(option.add_run(f"{letter}. "), size=10.2, color=NAVY, bold=True)
        set_run_font(option.add_run(alternative), size=10.2)

    correct_index = question["correct_answer_index"]
    correct_letter = letters[correct_index]
    answer = doc.add_table(rows=1, cols=1)
    answer.style = "Table Grid"
    set_table_geometry(answer, [9360])
    cell = answer.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, color="BFC7D4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(f"GABARITO EDITORIAL: {correct_letter} — "), size=10.2, color=NAVY, bold=True)
    set_run_font(p.add_run(question["alternatives"][correct_index]), size=10.2)

    curadoria = doc.add_paragraph()
    curadoria.paragraph_format.space_before = Pt(11)
    curadoria.paragraph_format.space_after = Pt(5)
    set_run_font(curadoria.add_run("CURADORIA  "), size=10.5, color=NAVY, bold=True)
    set_run_font(curadoria.add_run("☐ Aprovar     ☐ Revisar     ☐ Descartar"), size=10.5, bold=True)

    reviewer = doc.add_paragraph()
    reviewer.paragraph_format.space_after = Pt(6)
    set_run_font(reviewer.add_run("Especialista: ______________________________    Data: ____ / ____ / ______"), size=9.5, color=MUTED)

    observations = doc.add_table(rows=1, cols=1)
    observations.style = "Table Grid"
    set_table_geometry(observations, [9360])
    cell = observations.cell(0, 0)
    set_cell_border(cell, color="BFC7D4")
    row = observations.rows[0]
    row.height = Inches(0.95)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("OBSERVAÇÕES"), size=9.5, color=MUTED, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("Clique aqui e registre observações do especialista."), size=9.5, color="7A8491", italic=True)


def build_document() -> None:
    questions = load_questions()
    selected = select_questions(questions)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    style_document(doc)
    add_title_page(doc, selected)
    for topic, topic_questions in selected.items():
        for index, question in enumerate(topic_questions, start=1):
            add_question_form(doc, question, topic, index)

    doc.core_properties.title = "Caderno de Curadoria - Questões da Fase 3"
    doc.core_properties.subject = "Seleção das melhores questões por tópico para avaliação especializada"
    doc.core_properties.author = "AIIMS"
    doc.core_properties.comments = f"Gerado em {date.today().isoformat()} a partir de questoes_fase3.jsonl"
    doc.save(OUTPUT_PATH)
    print(f"Documento criado: {OUTPUT_PATH}")
    print(f"Tópicos: {len(selected)} | Questões: {sum(len(items) for items in selected.values())}")
    for topic, topic_questions in selected.items():
        print(f"- {topic}: {[question['id'] for question in topic_questions]}")


if __name__ == "__main__":
    build_document()
